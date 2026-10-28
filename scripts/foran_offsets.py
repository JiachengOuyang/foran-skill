# -*- coding: utf-8 -*-
"""Parse a FSURF loftbook (.lis offset table) and integrate hydrostatics.

The `File > Ship > Offset table` command writes a "LOFTBOOK OF HULL FORMS"
report; chapter 1 has the main particulars and chapter 3 the half-breadths in
frames and waterlines (millimetres).  From chapter 3 this module computes:

  * section areas and displacement volume at given draughts
  * block coefficient Cb (Lpp / moulded breadth / draught)
  * LCB measured from A.P. (positive forward)
  * midship section coefficient Cm and waterplane coefficient Cw

All integration is plain trapezoidal; the report rounds half-breadths to 1 mm,
so expect ~0.1% agreement with FORAN/FBASIC.

Usage:
  python foran_offsets.py parse --lis offsets.lis
  python foran_offsets.py parse --lis offsets.lis --draughts 9.83,10.27 --json
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path


def _floats(text: str) -> list[float]:
    return [float(x) for x in re.findall(r"-?\d+(?:\.\d+)?", text)]


def parse_loftbook(path: Path) -> dict:
    """Return {'lpp':..., 'breadth':..., 'draught':..., 'frames': {no: (x, {z: y_m})}}."""
    lines = Path(path).read_text(encoding="utf-8", errors="replace").splitlines()
    lpp = breadth = draught = None
    frames: dict[float, dict] = {}
    cur_frames: list[float] = []
    cur_dists: dict[float, float] = {}
    for i, line in enumerate(lines):
        s = line.strip()
        if s.startswith("LENGTH BETWEEN PERPENDICULARS"):
            v = _floats(s)
            if v:
                lpp = v[-1] if v[-1] > 1 else v[0]
        elif s.startswith("MOULDED BREADTH"):
            v = _floats(s)
            if v:
                breadth = v[-1] if v[-1] > 1 else v[0]
        elif s.startswith("MOULDED DRAUGHT"):
            v = _floats(s)
            if v:
                draught = v[-1]
        elif "END OF CHAPTER 3" in s or "CHAPTER 4" in s:
            cur_frames = []
        elif s.startswith("FRAME NO."):
            nums = _floats(s[len("FRAME NO."):])
            cur_frames = nums
            if i + 1 < len(lines) and lines[i + 1].strip().startswith("DIST A.P."):
                d = _floats(lines[i + 1].strip()[len("DIST A.P."):])
                if len(d) >= len(cur_frames):
                    cur_dists = {f: x for f, x in zip(cur_frames, d)}
            continue
        elif s.startswith("DIST A.P."):
            d = _floats(s[len("DIST A.P."):])
            if cur_frames and len(d) >= len(cur_frames):
                cur_dists = {f: x for f, x in zip(cur_frames, d)}
            continue
        if not cur_frames:
            continue
        m = re.match(r"^(\d+(?:\.\d+)?)\s+((?:-?\d+\s*)+)$", s)
        if not m:
            continue
        z = float(m.group(1))
        vals = [int(v) for v in m.group(2).split()]
        for f, v in zip(cur_frames, vals):
            x = cur_dists.get(f)
            if x is None:
                continue
            entry = frames.setdefault(f, {"x": x, "z": {}})
            entry["z"][z] = v / 1000.0
    return {"lpp": lpp, "breadth": breadth, "draught": draught, "frames": frames}


def repair_internal_gaps(zmap: dict[float, float],
                         max_gap: float = 3.0) -> tuple[dict, list]:
    """Fill zero runs that are bounded by nonzero half-breadths.

    FSURF's loftbook can report 0 at a station that grazes a patch seam or a
    vertical tangency even though the hull is continuous (the Surface
    Validator still reports no free edges).  Such internal zeros make the
    section area collapse; interpolate across short runs only.
    """
    pts = sorted(zmap.items())
    out = dict(pts)
    repairs = []
    i = 1
    while i < len(pts) - 1:
        z, y = pts[i]
        if y != 0:
            i += 1
            continue
        j = i
        while j < len(pts) and pts[j][1] == 0:
            j += 1
        if j >= len(pts) or pts[i - 1][1] <= 0:
            i = j
            continue
        z0, y0 = pts[i - 1]
        z1, y1 = pts[j]
        if z1 - z0 <= max_gap:
            for k in range(i, j):
                zk = pts[k][0]
                out[zk] = (y0 + (y1 - y0) * (zk - z0) / (z1 - z0)
                           if z1 > z0 else y0)
            repairs.append((z0, z1))
        i = j
    return out, repairs


def repair_frames(frames: dict, max_gap: float = 3.0) -> list:
    """Repair FSURF loftbook seam zeros using the neighbouring stations.

    At a patch seam the loftbook may report a whole band of zeros at one
    station (e.g. frames 14/15 of the SAMPLE model below z=2 m) although the
    surface is continuous.  Interpolating along z from the two bounding
    points underestimates a convex bilge; interpolating the same waterline
    between the neighbouring stations keeps the true shape.
    """
    order = sorted(frames, key=lambda f: frames[f]["x"])

    for idx, f in enumerate(order):
        zmap = frames[f]["z"]
        pts = sorted(zmap.items())
        i = 1
        while i < len(pts) - 1:
            z, y = pts[i]
            if y != 0:
                i += 1
                continue
            j = i
            while j < len(pts) and pts[j][1] == 0:
                j += 1
            if j >= len(pts) or pts[i - 1][1] <= 0:
                i = j
                continue
            if pts[j][0] - pts[i - 1][0] <= max_gap:
                for k in range(i, j):
                    zk = pts[k][0]
                    lv = rv = None
                    for ii in range(idx - 1, -1, -1):
                        v = frames[order[ii]]["z"].get(zk)
                        if v and v > 0:
                            lv = (frames[order[ii]]["x"], v)
                            break
                    for ii in range(idx + 1, len(order)):
                        v = frames[order[ii]]["z"].get(zk)
                        if v and v > 0:
                            rv = (frames[order[ii]]["x"], v)
                            break
                    if lv and rv and rv[0] > lv[0]:
                        x = frames[f]["x"]
                        zmap[zk] = lv[1] + (rv[1] - lv[1]) * (x - lv[0]) / (
                            rv[0] - lv[0])
            i = j

    # Any remaining internal run (no usable neighbours) is filled along z.
    for f in order:
        fixed, _rep = repair_internal_gaps(frames[f]["z"], max_gap)
        frames[f]["z"] = fixed
    return order


def _interp(zmap: dict[float, float], z: float) -> float:
    pts = sorted(zmap.items())
    if not pts:
        return 0.0
    if z <= pts[0][0]:
        return pts[0][1]
    for (z0, y0), (z1, y1) in zip(pts, pts[1:]):
        if z0 <= z <= z1:
            return y0 + (y1 - y0) * (z - z0) / (z1 - z0) if z1 > z0 else y1
    return pts[-1][1]


def _section_area(zmap: dict[float, float], draught: float) -> float:
    """Half-breadth integral from 0 to `draught`, times 2 (full breadth)."""
    pts = sorted(zmap.items())
    if not pts:
        return 0.0
    prev_z = 0.0
    prev_y = pts[0][1]
    half = 0.0
    for z, y in pts:
        if z <= 0:
            continue
        if z >= draught:
            yT = _interp(zmap, draught)
            if z > prev_z:
                half += (prev_y + yT) / 2.0 * max(0.0, draught - prev_z)
            return 2.0 * half
        half += (prev_y + y) / 2.0 * (z - prev_z)
        prev_z, prev_y = z, y
    return 2.0 * half


def _trapz(points: list[tuple[float, float]]) -> float:
    return sum((x1 - x0) * (y0 + y1) / 2.0
               for (x0, y0), (x1, y1) in zip(points, points[1:]))


def hydrostatics(data: dict, draughts: list[float],
                 lpp: float | None = None,
                 breadth: float | None = None,
                 repair: bool = True) -> list[dict]:
    lpp = lpp or data.get("lpp")
    breadth = breadth or data.get("breadth")
    frames = data["frames"]
    if not lpp or not breadth:
        raise SystemExit("Lpp/breadth not found in the report; pass --lpp/--breadth")
    if repair:
        repair_frames(frames)
    xs = sorted((v["x"], v["z"]) for v in frames.values())
    out = []
    for t in draughts:
        sections = [(x, _section_area(z, t)) for x, z in xs]
        disp = _trapz(sections)
        moment = _trapz([(x, x * a) for x, a in sections])
        lcb = moment / disp if disp else 0.0
        cm = max((a for _x, a in sections), default=0.0) / (breadth * t)
        wp = _trapz([(x, 2.0 * _interp(z, t)) for x, z in xs])
        out.append({
            "draught": t,
            "volume": disp,
            "cb": disp / (lpp * breadth * t) if t else 0.0,
            "lcb_ap": lcb,
            "cm": cm,
            "cwp": wp / (lpp * breadth),
            "waterplane": wp,
        })
    return out


def parse_expect(spec: str) -> dict:
    """Parse `9.83:42514,10.27:44644` or `T:volume:cb` into {T: (vol, cb)}."""
    out = {}
    for part in spec.split(","):
        part = part.strip()
        if not part:
            continue
        bits = part.split(":")
        out[float(bits[0])] = (float(bits[1]),
                               float(bits[2]) if len(bits) > 2 else None)
    return out


def compare(rows: list[dict], spec: str,
            tol: float = 0.005) -> tuple[list[str], bool]:
    """Compare integrated rows with targets (relative for vol, absolute for Cb)."""
    targets = parse_expect(spec)
    lines = []
    ok = True
    for r in rows:
        t = r["draught"]
        tgt = next((v for k, v in targets.items() if abs(k - t) < 0.01), None)
        if tgt is None:
            continue
        vol, cb = tgt
        dv = (r["volume"] - vol) / vol if vol else 0.0
        good = abs(dv) <= tol
        msg = (f"T={t:6.2f}  vol={r['volume']:10.1f}  expect {vol:10.1f}  "
               f"Δ={dv * 100:+.2f}%")
        if cb is not None:
            dc = r["cb"] - cb
            good = good and abs(dc) <= tol
            msg += f"  Cb={r['cb']:.4f}  expect {cb:.4f}  Δ={dc:+.4f}"
        msg += "  PASS" if good else "  FAIL"
        lines.append(msg)
        ok = ok and good
    return lines, ok


def _interp_y(zmap: dict, z: float) -> float:
    """Half-breadth at waterline z (linear between the reported levels)."""
    pts = sorted(zmap.items())
    if not pts:
        return 0.0
    if z <= pts[0][0]:
        # The first loftbook level is ~1 mm above the keel; for a flat bottom
        # the half-breadth at z=0 equals that first value, not zero.
        return pts[0][1] if z >= 0.0 else 0.0
    if z >= pts[-1][0]:
        return pts[-1][1]
    for (z0, y0), (z1, y1) in zip(pts, pts[1:]):
        if z0 <= z <= z1:
            if z1 == z0:
                return y1
            return y0 + (y1 - y0) * (z - z0) / (z1 - z0)
    return 0.0


def section_area(zmap: dict, draught: float) -> float:
    """Full section area (m^2) from the keel to `draught`."""
    pts = [(z, y) for z, y in sorted(zmap.items()) if z < draught - 1e-9]
    zs = [0.0] + [z for z, _ in pts] + [draught]
    ys = [_interp_y(zmap, z) for z in zs]
    area = 0.0
    for i in range(1, len(zs)):
        area += 0.5 * (ys[i] + ys[i - 1]) * (zs[i] - zs[i - 1])
    return 2.0 * area


def export_tables(data: dict, prefix: Path, draught: float,
                  wl_step: float = 0.5, max_wl: float = 15.0,
                  make_xlsx: bool = True, title: str = "横剖面面积曲线",
                  stations: int = 0) -> dict:
    """Write the offset table (CSV/XLSX) and the section-area curve."""
    frames = data["frames"]
    order = repair_frames(frames)
    levels = []
    k = 1
    while wl_step * k <= max_wl + 1e-9:
        levels.append(round(wl_step * k, 6))
        k += 1
    if not any(abs(lv - draught) < 1e-6 for lv in levels):
        levels.append(draught)
        levels.sort()

    frame_x = [frames[f]["x"] for f in order]
    frame_y = [[_interp_y(frames[f]["z"], lv) for lv in levels] for f in order]

    def at_x(x):
        """Interpolate the half-breadths between the two bracketing frames."""
        if x < frame_x[0] or x > frame_x[-1]:
            return [0.0] * len(levels)
        for i in range(len(frame_x) - 1):
            x0, x1 = frame_x[i], frame_x[i + 1]
            if x0 <= x <= x1:
                t = 0.0 if x1 == x0 else (x - x0) / (x1 - x0)
                return [y0 + (y1 - y0) * t
                        for y0, y1 in zip(frame_y[i], frame_y[i + 1])]
        return frame_y[-1]

    table_rows = []
    if stations and stations >= 2:
        lpp = data.get("lpp") or frame_x[-1]
        for i in range(stations):
            x = lpp * i / (stations - 1)
            table_rows.append([i, x] + at_x(x))
    else:
        for f in order:
            table_rows.append([f, frames[f]["x"]] + frame_y[order.index(f)])

    area_rows = []
    for f in order:
        area_rows.append([f, frames[f]["x"], section_area(frames[f]["z"], draught)])
    volume = 0.0
    for (f0, x0, a0), (f1, x1, a1) in zip(area_rows, area_rows[1:]):
        volume += 0.5 * (a0 + a1) * abs(x1 - x0)

    prefix = Path(prefix)
    prefix.parent.mkdir(parents=True, exist_ok=True)
    csv_table = prefix.with_name(prefix.name + "_型值表.csv")
    csv_area = prefix.with_name(prefix.name + "_面积曲线.csv")
    svg_area = prefix.with_name(prefix.name + "_面积曲线.svg")
    png_area = prefix.with_name(prefix.name + "_面积曲线.png")

    def _num(v):
        return f"{v:.4f}".rstrip("0").rstrip(".")

    with open(csv_table, "w", encoding="utf-8-sig", newline="") as fh:
        fh.write("站号,X(m)," + ",".join(f"WL{_num(lv)}m" for lv in levels) + "\n")
        for row in table_rows:
            fh.write(f"{_num(row[0])},{row[1]:.4f},"
                     + ",".join(f"{v:.3f}" for v in row[2:]) + "\n")
    with open(csv_area, "w", encoding="utf-8-sig", newline="") as fh:
        fh.write("站号,X(m),面积(m2)\n")
        for f, x, a in area_rows:
            fh.write(f"{_num(f)},{x:.4f},{a:.4f}\n")

    xlsx = None
    if make_xlsx:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font, Alignment
        except Exception as exc:  # noqa: BLE001
            print(f"offsets export: openpyxl missing ({exc}); CSV only",
                  file=sys.stderr)
        else:
            xlsx = prefix.with_name(prefix.name + "_型值表.xlsx")
            wb = Workbook()
            ws = wb.active
            ws.title = "型值表"
            ws.append(["型值表（半宽，m）"])
            ws["A1"].font = Font(bold=True, size=12)
            ws.append(["站号", "X(m)"] + [f"WL {_num(lv)}" for lv in levels])
            for c in ws[2]:
                c.font = Font(bold=True)
                c.alignment = Alignment(horizontal="center")
            for row in table_rows:
                ws.append([row[0], row[1]] + [round(v, 3) for v in row[2:]])
            for i, lv in enumerate(levels, 3):
                ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = 9
            ws.column_dimensions["A"].width = 7
            ws.column_dimensions["B"].width = 10
            ws.freeze_panes = "C3"

            ws2 = wb.create_sheet("横剖面面积")
            ws2.append([f"横剖面面积  T={draught:g} m  ∫A dx={volume:,.1f} m3"])
            ws2["A1"].font = Font(bold=True, size=12)
            ws2.append(["站号", "X(m)", "面积(m²)"])
            for c in ws2[2]:
                c.font = Font(bold=True)
            for f, x, a in area_rows:
                ws2.append([f, x, round(a, 4)])
            ws2.column_dimensions["A"].width = 8
            ws2.column_dimensions["B"].width = 10
            ws2.column_dimensions["C"].width = 12
            wb.save(xlsx)

    _plot_area(area_rows, svg_area, png_area, title, draught)
    return {"table_csv": csv_table, "table_xlsx": xlsx, "area_csv": csv_area,
            "area_svg": svg_area, "area_png": png_area,
            "volume": volume, "draught": draught, "levels": levels,
            "rows": len(table_rows)}


def _plot_area(rows, svg_path: Path, png_path: Path, title: str,
               draught: float) -> None:
    """Small dependency-free SVG/PNG plot of the section-area curve."""
    if not rows:
        return
    xs = [r[1] for r in rows]
    ys = [r[2] for r in rows]
    x0, x1 = min(xs), max(xs)
    y1 = max(ys) or 1.0
    W, H, L, R, T, B = 1200, 760, 90, 40, 70, 80
    pw, ph = W - L - R, H - T - B

    def px(x):
        return L + (x - x0) / (x1 - x0 or 1) * pw

    def py(y):
        return T + ph - y / (y1 * 1.08) * ph

    pts = [(px(x), py(y)) for x, y in zip(xs, ys)]
    poly = " ".join(f"{a:.1f},{b:.1f}" for a, b in pts)
    grid = []
    for i in range(6):
        gx = L + pw * i / 5
        gv = x0 + (x1 - x0) * i / 5
        grid.append(f'<line x1="{gx:.1f}" y1="{T}" x2="{gx:.1f}" y2="{T+ph}" '
                    f'stroke="#dddddd"/>')
        grid.append(f'<text x="{gx:.1f}" y="{T+ph+24}" font-size="16" '
                    f'text-anchor="middle">{gv:.0f}</text>')
    for i in range(6):
        gy = T + ph - ph * i / 5
        gv = y1 * i / 5
        grid.append(f'<line x1="{L}" y1="{gy:.1f}" x2="{L+pw}" y2="{gy:.1f}" '
                    f'stroke="#dddddd"/>')
        grid.append(f'<text x="{L-10}" y="{gy+5:.1f}" font-size="16" '
                    f'text-anchor="end">{gv:.0f}</text>')
    svg = f'''<svg xmlns="http://www.w3.org/2000/svg" width="{W}" height="{H}">
<rect width="100%" height="100%" fill="white"/>
<text x="{W/2:.0f}" y="38" font-size="26" text-anchor="middle"
 font-family="Microsoft YaHei, SimHei, sans-serif">{title}（T={draught:g} m）</text>
{''.join(grid)}
<line x1="{L}" y1="{T}" x2="{L}" y2="{T+ph}" stroke="black"/>
<line x1="{L}" y1="{T+ph}" x2="{L+pw}" y2="{T+ph}" stroke="black"/>
<polyline points="{poly}" fill="none" stroke="#1f4e79" stroke-width="2.5"/>
<text x="{L+pw/2:.0f}" y="{H-30}" font-size="20" text-anchor="middle"
 font-family="Microsoft YaHei, SimHei, sans-serif">X（自 AP，m）</text>
<text x="24" y="{T+ph/2:.0f}" font-size="20" text-anchor="middle"
 transform="rotate(-90 24 {T+ph/2:.0f})"
 font-family="Microsoft YaHei, SimHei, sans-serif">横剖面面积 A（m²）</text>
</svg>'''
    svg_path.write_text(svg, encoding="utf-8")

    try:
        from PIL import Image, ImageDraw, ImageFont
    except Exception:  # noqa: BLE001
        return

    def font(size):
        for name in (r"C:\Windows\Fonts\msyh.ttc", r"C:\Windows\Fonts\simhei.ttf",
                     r"C:\Windows\Fonts\arial.ttf"):
            try:
                return ImageFont.truetype(name, size)
            except Exception:  # noqa: BLE001
                continue
        return ImageFont.load_default()

    img = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(img)
    f28, f20, f16 = font(28), font(20), font(16)
    tw = d.textbbox((0, 0), f"{title}（T={draught:g} m）", font=f28)
    d.text(((W - (tw[2] - tw[0])) / 2, 24), f"{title}（T={draught:g} m）",
           fill="black", font=f28)
    for i in range(6):
        gx = L + pw * i / 5
        gv = x0 + (x1 - x0) * i / 5
        d.line([(gx, T), (gx, T + ph)], fill="#dddddd")
        d.text((gx - 12, T + ph + 8), f"{gv:.0f}", fill="black", font=f16)
    for i in range(6):
        gy = T + ph - ph * i / 5
        gv = y1 * i / 5
        d.line([(L, gy), (L + pw, gy)], fill="#dddddd")
        d.text((L - 50, gy - 8), f"{gv:.0f}", fill="black", font=f16)
    d.line([(L, T), (L, T + ph)], fill="black", width=2)
    d.line([(L, T + ph), (L + pw, T + ph)], fill="black", width=2)
    d.line(pts, fill="#1f4e79", width=3)
    d.text((L + pw / 2 - 70, H - 48), "X（自 AP，m）", fill="black", font=f20)
    label = Image.new("RGBA", (260, 32), (255, 255, 255, 0))
    ld = ImageDraw.Draw(label)
    ld.text((0, 0), "横剖面面积 A（m²）", fill="black", font=f20)
    label = label.rotate(90, expand=True)
    img.paste(label, (10, int(T + ph / 2 - 130)), label)
    img.save(png_path)


def _cmd_export(a) -> int:
    data = parse_loftbook(Path(a.lis).expanduser().resolve())
    draught = a.draught if a.draught is not None else (data.get("draught") or 9.83)
    res = export_tables(data, Path(a.out_prefix), draught,
                        wl_step=a.waterline_step, max_wl=a.max_waterline,
                        make_xlsx=not a.no_xlsx, title=a.title,
                        stations=a.stations)
    print(f"rows     : {res['rows']} stations, {len(res['levels'])} waterlines, "
          f"T={draught:g} m")
    print(f"volume   : {res['volume']:,.1f} m3 (∫A dx, trapezoidal)")
    for key in ("table_csv", "table_xlsx", "area_csv", "area_svg", "area_png"):
        p = res.get(key)
        if p:
            print(f"{key:9s}: {p} ({Path(p).stat().st_size} bytes)")
    return 0


def main(argv=None) -> int:
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
    except Exception:  # noqa: BLE001
        pass
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("parse", help="integrate hydrostatics from a .lis loftbook")
    p.add_argument("--lis", required=True)
    p.add_argument("--draughts", default="9.83,10.27")
    p.add_argument("--lpp", type=float)
    p.add_argument("--breadth", type=float)
    p.add_argument("--json", action="store_true")
    p.add_argument("--no-repair", action="store_true",
                   help="do not interpolate internal zero runs in the report")
    p.add_argument("--expect",
                   help="targets T:volume[:Cb], e.g. 9.83:42514,10.27:44644")
    p.add_argument("--tol", type=float, default=0.005,
                   help="relative vol / absolute Cb tolerance (default 0.005)")
    p.set_defaults(fn=_cmd_parse)

    p = sub.add_parser("export",
                       help="write 型值表 (CSV/XLSX) + section-area curve")
    p.add_argument("--lis", required=True)
    p.add_argument("--out-prefix", required=True,
                   help="output path prefix, e.g. out/SAMPLE")
    p.add_argument("--draught", type=float,
                   help="draught for the area curve (default: file value)")
    p.add_argument("--waterline-step", type=float, default=0.5)
    p.add_argument("--max-waterline", type=float, default=15.0)
    p.add_argument("--no-xlsx", action="store_true")
    p.add_argument("--stations", type=int, default=0,
                   help="resample the table at N equally spaced stations "
                        "(e.g. 21); 0 = keep the loftbook frames")
    p.add_argument("--title", default="横剖面面积曲线")
    p.set_defaults(fn=_cmd_export)
    a = ap.parse_args(argv)
    return a.fn(a)


def _cmd_parse(a) -> int:
    data = parse_loftbook(Path(a.lis).expanduser().resolve())
    ts = [float(x) for x in a.draughts.split(",") if x.strip()]
    rows = hydrostatics(data, ts, a.lpp, a.breadth, repair=not a.no_repair)
    if a.json:
        print(json.dumps({"lpp": data["lpp"], "breadth": data["breadth"],
                          "results": rows}, ensure_ascii=False, indent=2))
        return 0
    print(f"lpp={data['lpp']}  breadth={data['breadth']}  "
          f"frames={len(data['frames'])}")
    for r in rows:
        print(f"T={r['draught']:6.2f} m  vol={r['volume']:10.1f} m3  "
              f"Cb={r['cb']:.4f}  LCB(AP)={r['lcb_ap']:7.2f} m  "
              f"Cm={r['cm']:.4f}  Cw={r['cwp']:.4f}")
    if a.expect:
        lines, ok = compare(rows, a.expect, a.tol)
        for ln in lines:
            print(ln)
        print("check    : " + ("PASS" if ok else "FAIL"))
        return 0 if ok else 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
